"""Arquivos sintéticos com dados fabricados — nunca uma fatura real."""

import datetime

import openpyxl
import pytest

from app_fi.data.statement_files import read_csv, read_statement, read_xlsx

_HEADER = [
    "Data de Compra", "Nome no Cartão", "Final do Cartão", "Categoria",
    "Descrição", "Parcela", "Valor (em US$)", "Cotação (em R$)", "Valor (em R$)",
]


def test_read_csv_parses_header_and_rows(tmp_path):
    path = tmp_path / "fatura.csv"
    path.write_text(
        ";".join(_HEADER) + "\n"
        "10/08/2026;FULANO DE TAL;0000;Educacional;LOJA TESTE;Única;0;0;99.90\n",
        encoding="utf-8-sig",
    )
    rows = read_csv(path)
    assert len(rows) == 1
    assert rows[0]["Descrição"] == "LOJA TESTE"
    assert rows[0]["Valor (em R$)"] == "99.90"


def test_read_csv_falls_back_to_cp1252(tmp_path):
    # exportação de banco em Windows-1252/Latin-1 (comum), não UTF-8
    path = tmp_path / "fatura.csv"
    conteudo = ";".join(_HEADER) + "\n" + (
        "10/08/2026;FULANO DE TAL;0000;Educacional;LOJA TESTE ção;Única;0;0;99.90\n"
    )
    path.write_bytes(conteudo.encode("cp1252"))

    rows = read_csv(path)
    assert rows[0]["Descrição"] == "LOJA TESTE ção"


def test_read_csv_raises_readable_error_for_unsupported_encoding(tmp_path):
    path = tmp_path / "fatura.csv"
    # 0x81 não existe nem em utf-8 nem em cp1252 (cp1252 mapeia quase todo byte,
    # então um arquivo aleatório raramente falha nos dois ao mesmo tempo)
    path.write_bytes(b"Data de Compra\x81\xff")

    with pytest.raises(ValueError):
        read_csv(path)


def test_read_xlsx_parses_header_and_rows(tmp_path):
    path = tmp_path / "fatura.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(_HEADER)
    ws.append(["10/08/2026", "FULANO DE TAL", "0000", "Educacional", "LOJA TESTE", "Única", 0, 0, 99.90])
    wb.save(path)

    rows = read_xlsx(path)
    assert len(rows) == 1
    assert rows[0]["Descrição"] == "LOJA TESTE"
    assert rows[0]["Valor (em R$)"] == "99.9"


def test_read_xlsx_normalizes_date_cell_to_ddmmyyyy(tmp_path):
    path = tmp_path / "fatura.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(_HEADER)
    ws.append([datetime.date(2026, 8, 10), "FULANO DE TAL", "0000", "-", "LOJA", "Única", 0, 0, 10])
    wb.save(path)

    rows = read_xlsx(path)
    assert rows[0]["Data de Compra"] == "10/08/2026"


def test_read_xlsx_skips_fully_blank_rows(tmp_path):
    path = tmp_path / "fatura.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(_HEADER)
    ws.append(["10/08/2026", "FULANO DE TAL", "0000", "-", "LOJA", "Única", 0, 0, 10])
    ws.append([None] * len(_HEADER))
    wb.save(path)

    rows = read_xlsx(path)
    assert len(rows) == 1


def test_read_statement_dispatches_by_extension(tmp_path):
    csv_path = tmp_path / "a.csv"
    csv_path.write_text(";".join(_HEADER) + "\n", encoding="utf-8-sig")
    xlsx_path = tmp_path / "a.xlsx"
    wb = openpyxl.Workbook()
    wb.active.append(_HEADER)
    wb.save(xlsx_path)

    assert read_statement(csv_path) == read_csv(csv_path)
    assert read_statement(xlsx_path) == read_xlsx(xlsx_path)
